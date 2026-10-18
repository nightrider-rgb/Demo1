"""
Folder Lookup - Windows desktop app (Python + Tkinter)

Pick a folder (or files), read .txt / .csv / .tsv / Excel files, choose to read
ALL files or only SELECTED files, then find rows by a header ID column or by any
value anywhere in the data. IDs can be pasted or bulk-uploaded from a list file.
Results can be exported to CSV or Excel.

SETUP (once, in Command Prompt or PowerShell):
    pip install openpyxl xlrd

RUN:
    python folder_lookup.py

OPTIONAL - build a single .exe you can double-click:
    pip install pyinstaller
    pyinstaller --onefile --windowed folder_lookup.py
    (the exe appears in the "dist" folder)

Notes:
    - openpyxl reads .xlsx / .xlsm and writes Excel exports.
    - xlrd is only needed for old .xls files.
    - Everything runs locally. Nothing is uploaded anywhere.
"""
from __future__ import annotations

import csv
import io
import os
import queue
import re
import threading
import tkinter as tk
from dataclasses import dataclass, field
from tkinter import filedialog, messagebox, ttk
from typing import List, Optional

try:
    import openpyxl
except ImportError:  # handled when an Excel file is read
    openpyxl = None
try:
    import xlrd
except ImportError:
    xlrd = None

csv.field_size_limit(2**31 - 1)

TEXT_EXT = {".txt", ".csv", ".tsv"}
EXCEL_EXT = {".xlsx", ".xlsm", ".xls"}
SUPPORTED = TEXT_EXT | EXCEL_EXT
MAX_SHOWN = 5000  # rows drawn in the results table (exports include everything)
ILLEGAL_XLSX = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


# --------------------------------------------------------------------------
# Data model
# --------------------------------------------------------------------------
@dataclass
class Sheet:
    name: Optional[str]
    headers: List[str]
    rows: List[List[str]]


@dataclass
class Entry:
    path: str
    rel: str
    ext: str
    iid: str = ""
    selected: bool = False
    data: Optional[List[Sheet]] = None
    error: Optional[str] = None
    mtime: float = 0.0

    @property
    def is_excel(self) -> bool:
        return self.ext in EXCEL_EXT


# --------------------------------------------------------------------------
# Reading files
# --------------------------------------------------------------------------
def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_text(path: str) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with open(path, "r", encoding="latin-1") as f:
        return f.read()


def detect_delim(line: str) -> Optional[str]:
    counts = {"\t": 0, ",": 0, "|": 0, ";": 0}
    in_q = False
    for ch in line:
        if ch == '"':
            in_q = not in_q
        elif not in_q and ch in counts:
            counts[ch] += 1
    best, n = None, 0
    for k, c in counts.items():
        if c > n:
            best, n = k, c
    return best


def text_grid(path: str):
    """Return (grid, headerless). Headerless means a plain one-value-per-line file."""
    text = read_text(path)
    lines = [l for l in text.splitlines() if l.strip()]
    if not lines:
        return [], False
    d = detect_delim(lines[0])
    if d is None:
        return [[l.strip()] for l in lines], True
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), delimiter=d)]
    return rows, False


def excel_grids(path: str, ext: str, first_only: bool = False):
    """Yield (sheet_name, grid) for each sheet in an Excel file."""
    if ext == ".xls":
        if xlrd is None:
            raise RuntimeError("Reading .xls files needs:  pip install xlrd")
        wb = xlrd.open_workbook(path)
        for ws in wb.sheets():
            grid = [[cell_str(ws.cell_value(r, c)) for c in range(ws.ncols)] for r in range(ws.nrows)]
            yield ws.name, grid
            if first_only:
                return
    else:
        if openpyxl is None:
            raise RuntimeError("Reading Excel files needs:  pip install openpyxl")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                grid = [[cell_str(v) for v in row] for row in ws.iter_rows(values_only=True)]
                yield ws.title, grid
                if first_only:
                    return
        finally:
            wb.close()


def build_sheet(name, grid, headerless) -> Optional[Sheet]:
    grid = [r for r in grid if any(str(c).strip() for c in r)]
    if not grid:
        return None
    if headerless:
        headers, rows = ["Line"], grid
    else:
        headers, rows = [str(c).strip() for c in grid[0]], grid[1:]
    width = max([len(headers)] + [len(r) for r in rows])
    seen, final = {}, []
    for i in range(width):
        h = headers[i] if i < len(headers) and headers[i] else f"Column {i + 1}"
        k = h.lower()
        if k in seen:
            seen[k] += 1
            h = f"{h}_{seen[k]}"
        else:
            seen[k] = 1
        final.append(h)
    return Sheet(name, final, rows)


def read_entry(e: Entry) -> None:
    try:
        mtime = os.path.getmtime(e.path)
    except OSError as ex:
        e.data, e.error = None, str(ex)
        return
    if e.data is not None and e.mtime == mtime:
        return  # already read and unchanged
    e.error, e.data, e.mtime = None, None, mtime
    try:
        if e.is_excel:
            sheets = []
            for name, grid in excel_grids(e.path, e.ext):
                s = build_sheet(name, grid, False)
                if s:
                    sheets.append(s)
            e.data = sheets
        else:
            grid, headerless = text_grid(e.path)
            s = build_sheet(None, grid, headerless)
            e.data = [s] if s else []
    except Exception as ex:  # keep going if one file is bad
        e.error, e.data = str(ex) or "Could not read file", None


def read_first_column(path: str, skip_header: bool) -> List[str]:
    ext = os.path.splitext(path)[1].lower()
    if ext in EXCEL_EXT:
        grid = next(excel_grids(path, ext, first_only=True), (None, []))[1]
    else:
        grid, _ = text_grid(path)
    col = [str(r[0]).strip() for r in grid if r and str(r[0]).strip()]
    if skip_header and col:
        col = col[1:]
    return list(dict.fromkeys(col))


def parse_terms(text: str) -> List[str]:
    return [t.strip() for t in re.split(r"[\n\r,;\t]+", text) if t.strip()]


# --------------------------------------------------------------------------
# Searching
# --------------------------------------------------------------------------
def search(entries, mode, id_key, terms, progress=None):
    """mode 'id'  : exact (case-insensitive) match on the header whose lowercase name == id_key
       mode 'any' : row matches if any cell contains any term"""
    lower = {}
    for t in terms:
        lower.setdefault(t.lower(), t)
    keys = list(lower)
    keyset = set(keys)
    found, groups, skipped = set(), [], []
    for i, e in enumerate(entries):
        for s in e.data or []:
            label = e.rel + (f" / {s.name}" if e.is_excel else "")
            col = -1
            if mode == "id":
                col = next((j for j, h in enumerate(s.headers) if h.lower() == id_key), -1)
                if col < 0:
                    skipped.append(label)
                    continue
            hits = []
            for n, r in enumerate(s.rows, 1):
                if mode == "id":
                    v = (r[col] if col < len(r) else "").strip().lower()
                    if v in keyset:
                        found.add(v)
                        hits.append((n, r, lower[v]))
                else:
                    joined = "\x01".join(r).lower()
                    m = [k for k in keys if k in joined]
                    if m:
                        found.update(m)
                        hits.append((n, r, " | ".join(lower[k] for k in m)))
            if hits:
                groups.append(dict(path=e.rel, sheet=s.name if e.is_excel else None,
                                   headers=s.headers, rows=hits))
        if progress:
            progress(i + 1, len(entries))
    return dict(groups=groups, skipped=skipped, term_count=len(keys),
                not_found=[lower[k] for k in keys if k not in found])


def export_table(groups):
    keys, labels = [], {}
    for g in groups:
        for h in g["headers"]:
            k = h.lower()
            if k not in labels:
                labels[k] = h
                keys.append(k)
    out = [["File", "Sheet", "Matched"] + [labels[k] for k in keys]]
    for g in groups:
        idx = [next((i for i, h in enumerate(g["headers"]) if h.lower() == k), -1) for k in keys]
        for _, cells, term in g["rows"]:
            out.append([g["path"], g["sheet"] or "", term] +
                       [cells[i] if 0 <= i < len(cells) else "" for i in idx])
    return out


# --------------------------------------------------------------------------
# GUI
# --------------------------------------------------------------------------
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Folder lookup")
        self.geometry("1280x800")
        self.minsize(1000, 660)

        self.entries: List[Entry] = []
        self.cache = {}
        self.by_iid = {}
        self.iid_counter = 0
        self.source_dir = None
        self.source_files = None
        self.loaded = False
        self.read_entries: List[Entry] = []
        self.bulk: List[str] = []
        self.bulk_path = None
        self.results = None
        self.idcol_map = {}
        self.phase = "Reading"
        self.q = queue.Queue()

        self.include_sub = tk.BooleanVar(value=True)
        self.scope = tk.StringVar(value="all")
        self.filter = tk.StringVar()
        self.mode = tk.StringVar(value="id")
        self.id_col = tk.StringVar()
        self.bulk_header = tk.BooleanVar(value=True)
        self.status = tk.StringVar()
        self.src_note = tk.StringVar(value="Nothing chosen yet.")
        self.msg = tk.StringVar()
        self.summary = tk.StringVar(value="Results will appear here.")
        self.bulk_text = tk.StringVar()

        self._build_ui()
        self.after(80, self._poll)

    # ---------------- UI construction ----------------
    def _build_ui(self):
        style = ttk.Style(self)
        for t in ("vista", "clam"):
            if t in style.theme_names():
                style.theme_use(t)
                break
        style.configure("Head.TLabel", font=("Segoe UI", 11, "bold"))
        style.configure("Muted.TLabel", foreground="#5B6B78")
        style.configure("Err.TLabel", foreground="#B3261E")

        root = ttk.Frame(self, padding=10)
        root.pack(fill="both", expand=True)
        root.columnconfigure(0, minsize=470)
        root.columnconfigure(1, weight=1)
        root.rowconfigure(0, weight=1)

        left = ttk.Frame(root)
        left.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        left.columnconfigure(0, weight=1)
        left.rowconfigure(1, weight=1)
        right = ttk.Frame(root)
        right.grid(row=0, column=1, sticky="nsew")
        right.columnconfigure(0, weight=1)
        right.rowconfigure(2, weight=1)

        # 1. Source
        f1 = ttk.LabelFrame(left, text=" 1. Choose a source ", padding=8)
        f1.grid(row=0, column=0, sticky="ew", pady=(0, 8))
        ttk.Button(f1, text="Choose folder…", command=self._choose_folder).grid(row=0, column=0, padx=(0, 6))
        ttk.Button(f1, text="Choose files…", command=self._choose_files).grid(row=0, column=1, padx=(0, 12))
        ttk.Checkbutton(f1, text="Include subfolders", variable=self.include_sub,
                        command=self._on_sub_toggle).grid(row=0, column=2)
        ttk.Label(f1, textvariable=self.src_note, style="Muted.TLabel", wraplength=430).grid(
            row=1, column=0, columnspan=3, sticky="w", pady=(6, 0))

        # 2. Files
        f2 = ttk.LabelFrame(left, text=" 2. Pick files to read ", padding=8)
        f2.grid(row=1, column=0, sticky="nsew", pady=(0, 8))
        f2.columnconfigure(0, weight=1)
        f2.rowconfigure(2, weight=1)
        rb = ttk.Frame(f2)
        rb.grid(row=0, column=0, sticky="w")
        self.rb_all = ttk.Radiobutton(rb, text="All files (0)", value="all", variable=self.scope,
                                      command=self._on_scope)
        self.rb_all.pack(side="left", padx=(0, 14))
        self.rb_sel = ttk.Radiobutton(rb, text="Selected files (0)", value="selected", variable=self.scope,
                                      command=self._on_scope)
        self.rb_sel.pack(side="left")
        ff = ttk.Frame(f2)
        ff.grid(row=1, column=0, sticky="ew", pady=6)
        ff.columnconfigure(1, weight=1)
        ttk.Label(ff, text="Filter").grid(row=0, column=0, padx=(0, 6))
        ent = ttk.Entry(ff, textvariable=self.filter)
        ent.grid(row=0, column=1, sticky="ew")
        self.filter.trace_add("write", lambda *_: self._on_filter())

        tf = ttk.Frame(f2)
        tf.grid(row=2, column=0, sticky="nsew")
        tf.columnconfigure(0, weight=1)
        tf.rowconfigure(0, weight=1)
        self.tree = ttk.Treeview(tf, columns=("info",), selectmode="extended", height=8)
        self.tree.heading("#0", text="File (click to select, Ctrl/Shift for several)")
        self.tree.heading("info", text="Size / rows")
        self.tree.column("#0", width=300, stretch=True)
        self.tree.column("info", width=90, anchor="e", stretch=False)
        sb = ttk.Scrollbar(tf, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        sb.grid(row=0, column=1, sticky="ns")
        self.tree.bind("<<TreeviewSelect>>", self._on_select)

        bf = ttk.Frame(f2)
        bf.grid(row=3, column=0, sticky="w", pady=6)
        ttk.Button(bf, text="Select listed", command=self._select_listed).pack(side="left", padx=(0, 6))
        ttk.Button(bf, text="Clear selection", command=self._clear_selection).pack(side="left")

        rf = ttk.Frame(f2)
        rf.grid(row=4, column=0, sticky="ew")
        rf.columnconfigure(1, weight=1)
        self.read_btn = ttk.Button(rf, text="Read files", command=self._read_clicked)
        self.read_btn.grid(row=0, column=0, padx=(0, 8))
        self.progress = ttk.Progressbar(rf, mode="determinate")
        self.progress.grid(row=0, column=1, sticky="ew")
        ttk.Label(f2, textvariable=self.status, style="Muted.TLabel", wraplength=430).grid(
            row=5, column=0, sticky="w", pady=(4, 0))

        # 3. Find
        f3 = ttk.LabelFrame(left, text=" 3. Find rows ", padding=8)
        f3.grid(row=2, column=0, sticky="ew")
        f3.columnconfigure(0, weight=1)
        mf = ttk.Frame(f3)
        mf.grid(row=0, column=0, sticky="w")
        ttk.Radiobutton(mf, text="By header ID", value="id", variable=self.mode,
                        command=self._on_mode).pack(side="left", padx=(0, 14))
        ttk.Radiobutton(mf, text="Anywhere in the data", value="any", variable=self.mode,
                        command=self._on_mode).pack(side="left")
        self.idrow = ttk.Frame(f3)
        self.idrow.grid(row=1, column=0, sticky="ew", pady=(8, 0))
        self.idrow.columnconfigure(1, weight=1)
        ttk.Label(self.idrow, text="ID header").grid(row=0, column=0, padx=(0, 6))
        self.idbox = ttk.Combobox(self.idrow, textvariable=self.id_col, state="readonly")
        self.idbox.grid(row=0, column=1, sticky="ew")
        self.terms_label = ttk.Label(f3, text="IDs to find (one per line, or separated by commas)")
        self.terms_label.grid(row=2, column=0, sticky="w", pady=(8, 2))
        tw = ttk.Frame(f3)
        tw.grid(row=3, column=0, sticky="ew")
        tw.columnconfigure(0, weight=1)
        self.terms = tk.Text(tw, height=5, wrap="word", font=("Segoe UI", 10))
        tsb = ttk.Scrollbar(tw, orient="vertical", command=self.terms.yview)
        self.terms.configure(yscrollcommand=tsb.set)
        self.terms.grid(row=0, column=0, sticky="ew")
        tsb.grid(row=0, column=1, sticky="ns")
        bk = ttk.Frame(f3)
        bk.grid(row=4, column=0, sticky="ew", pady=(8, 0))
        ttk.Button(bk, text="Upload list file…", command=self._choose_bulk).pack(side="left", padx=(0, 8))
        ttk.Checkbutton(bk, text="First row is a header", variable=self.bulk_header,
                        command=self._reload_bulk).pack(side="left")
        self.bulk_row = ttk.Frame(f3)
        self.bulk_row.grid(row=5, column=0, sticky="w", pady=(4, 0))
        ttk.Label(self.bulk_row, textvariable=self.bulk_text, style="Muted.TLabel").pack(side="left", padx=(0, 8))
        ttk.Button(self.bulk_row, text="Remove", command=self._clear_bulk).pack(side="left")
        self.bulk_row.grid_remove()
        sf = ttk.Frame(f3)
        sf.grid(row=6, column=0, sticky="ew", pady=(10, 0))
        self.search_btn = ttk.Button(sf, text="Search", command=self._search_clicked, state="disabled")
        self.search_btn.pack(side="left", padx=(0, 10))
        ttk.Label(sf, textvariable=self.msg, style="Err.TLabel", wraplength=330).pack(side="left")

        # Results
        ttk.Label(right, textvariable=self.summary, style="Head.TLabel").grid(row=0, column=0, sticky="w")
        ef = ttk.Frame(right)
        ef.grid(row=1, column=0, sticky="w", pady=6)
        self.csv_btn = ttk.Button(ef, text="Export CSV", command=self._export_csv, state="disabled")
        self.csv_btn.pack(side="left", padx=(0, 6))
        self.xlsx_btn = ttk.Button(ef, text="Export Excel", command=self._export_xlsx, state="disabled")
        self.xlsx_btn.pack(side="left")

        nb = ttk.Notebook(right)
        nb.grid(row=2, column=0, sticky="nsew")
        t1 = ttk.Frame(nb)
        t1.columnconfigure(0, weight=1)
        t1.rowconfigure(0, weight=1)
        self.rtree = ttk.Treeview(t1, show="headings", columns=())
        ry = ttk.Scrollbar(t1, orient="vertical", command=self.rtree.yview)
        rx = ttk.Scrollbar(t1, orient="horizontal", command=self.rtree.xview)
        self.rtree.configure(yscrollcommand=ry.set, xscrollcommand=rx.set)
        self.rtree.grid(row=0, column=0, sticky="nsew")
        ry.grid(row=0, column=1, sticky="ns")
        rx.grid(row=1, column=0, sticky="ew")
        nb.add(t1, text="Matches")
        self.nf_text = self._text_tab(nb, "Not found")
        self.skip_text = self._text_tab(nb, "Skipped files")

    def _text_tab(self, nb, title):
        fr = ttk.Frame(nb)
        fr.columnconfigure(0, weight=1)
        fr.rowconfigure(0, weight=1)
        t = tk.Text(fr, wrap="word", font=("Segoe UI", 10), state="disabled")
        sb = ttk.Scrollbar(fr, orient="vertical", command=t.yview)
        t.configure(yscrollcommand=sb.set)
        t.grid(row=0, column=0, sticky="nsew")
        sb.grid(row=0, column=1, sticky="ns")
        nb.add(fr, text=title)
        return t

    @staticmethod
    def _set_text(widget, text):
        widget.configure(state="normal")
        widget.delete("1.0", "end")
        widget.insert("1.0", text)
        widget.configure(state="disabled")

    # ---------------- background work ----------------
    def run_bg(self, work, done):
        def target():
            try:
                res = work(lambda i, n: self.q.put(("p", i, n)))
            except Exception as ex:  # surface errors on the UI thread
                self.q.put(("e", ex))
                return
            self.q.put(("d", res, done))
        threading.Thread(target=target, daemon=True).start()

    def _poll(self):
        try:
            while True:
                m = self.q.get_nowait()
                if m[0] == "p":
                    self.progress.configure(value=m[1])
                    self.status.set(f"{self.phase} {m[1]} of {m[2]}…")
                elif m[0] == "d":
                    m[2](m[1])
                else:
                    self.read_btn.state(["!disabled"])
                    self.search_btn.state(["!disabled" if self.loaded else "disabled"])
                    messagebox.showerror("Folder lookup", str(m[1]))
        except queue.Empty:
            pass
        self.after(80, self._poll)

    # ---------------- source ----------------
    def _choose_folder(self):
        d = filedialog.askdirectory(title="Choose a folder")
        if d:
            self.source_dir, self.source_files = os.path.normpath(d), None
            self._rescan()

    def _choose_files(self):
        ps = filedialog.askopenfilenames(
            title="Choose files",
            filetypes=[("Text and Excel", "*.txt *.csv *.tsv *.xlsx *.xlsm *.xls"), ("All files", "*.*")])
        if ps:
            self.source_dir, self.source_files = None, [os.path.normpath(p) for p in ps]
            self._rescan()

    def _on_sub_toggle(self):
        if self.source_dir:
            self._rescan()

    def _rescan(self):
        paths, ignored = [], 0
        if self.source_dir:
            for root, _dirs, files in os.walk(self.source_dir):
                for fn in files:
                    if fn.startswith("~$"):  # Excel lock files
                        continue
                    if os.path.splitext(fn)[1].lower() in SUPPORTED:
                        paths.append(os.path.join(root, fn))
                    else:
                        ignored += 1
                if not self.include_sub.get():
                    break
            base = self.source_dir
        else:
            for p in self.source_files or []:
                if os.path.splitext(p)[1].lower() in SUPPORTED:
                    paths.append(p)
                else:
                    ignored += 1
            base = None
        self.entries = []
        for p in sorted(paths, key=str.lower):
            e = self.cache.get(p)
            if e is None:
                self.iid_counter += 1
                e = Entry(path=p, rel=os.path.relpath(p, base) if base else os.path.basename(p),
                          ext=os.path.splitext(p)[1].lower(), iid=f"e{self.iid_counter}")
                self.cache[p] = e
                self.by_iid[e.iid] = e
            else:
                e.rel = os.path.relpath(p, base) if base else os.path.basename(p)
            self.entries.append(e)
        n = len(self.entries)
        if n:
            note = f"{n} readable file{'s' if n != 1 else ''} found"
            note += f", {ignored} other file{'s' if ignored != 1 else ''} ignored." if ignored else "."
        else:
            note = "No .txt, .csv, .tsv or Excel files found in that selection."
        self.src_note.set(note)
        self._invalidate()
        self._refresh_tree()

    # ---------------- file list ----------------
    def _visible(self):
        q = self.filter.get().strip().lower()
        return [e for e in self.entries if not q or q in e.rel.lower()]

    def _scope_entries(self):
        v = self._visible()
        return v if self.scope.get() == "all" else [e for e in v if e.selected]

    def _refresh_tree(self):
        self.tree.delete(*self.tree.get_children())
        sel = []
        for e in self._visible():
            if e.error:
                info = "error"
            elif e.data is not None:
                info = f"{sum(len(s.rows) for s in e.data):,} rows"
            else:
                try:
                    kb = os.path.getsize(e.path) / 1024
                    info = f"{kb:,.0f} KB" if kb < 1024 else f"{kb / 1024:.1f} MB"
                except OSError:
                    info = ""
            self.tree.insert("", "end", iid=e.iid, text=e.rel, values=(info,))
            if e.selected:
                sel.append(e.iid)
        if sel:
            self.tree.selection_set(sel)
        self._update_counts()

    def _update_counts(self):
        v = self._visible()
        self.rb_all.configure(text=f"All files ({len(v)})")
        self.rb_sel.configure(text=f"Selected files ({sum(1 for e in v if e.selected)})")

    def _invalidate(self):
        if self.loaded:
            self.status.set("Selection changed. Click Read files again.")
        self.loaded = False
        self.search_btn.state(["disabled"])

    def _on_select(self, _evt):
        sel = set(self.tree.selection())
        changed = False
        for iid in self.tree.get_children():
            e = self.by_iid[iid]
            new = iid in sel
            if e.selected != new:
                e.selected, changed = new, True
        if changed:
            if sel and self.scope.get() == "all":
                self.scope.set("selected")
            self._invalidate()
            self._update_counts()

    def _on_scope(self):
        self._invalidate()

    def _on_filter(self):
        self._invalidate()
        self._refresh_tree()

    def _select_listed(self):
        kids = self.tree.get_children()
        if kids:
            self.tree.selection_set(kids)

    def _clear_selection(self):
        for e in self.entries:
            e.selected = False
        self._invalidate()
        self._refresh_tree()

    # ---------------- reading ----------------
    def _read_clicked(self):
        lst = self._scope_entries()
        if not lst:
            messagebox.showinfo("Folder lookup", "Choose a source and pick at least one file first.")
            return
        self.phase = "Reading"
        self.read_btn.state(["disabled"])
        self.search_btn.state(["disabled"])
        self.progress.configure(maximum=len(lst), value=0)
        self.status.set(f"Reading 0 of {len(lst)}…")

        def work(progress):
            for i, e in enumerate(lst):
                read_entry(e)
                progress(i + 1, len(lst))
            return lst

        self.run_bg(work, self._read_done)

    def _read_done(self, lst):
        self.read_entries = [e for e in lst if e.data is not None]
        failed = len(lst) - len(self.read_entries)
        rows = sum(len(s.rows) for e in self.read_entries for s in e.data)
        self._build_header_options()
        self.loaded = True
        self.read_btn.state(["!disabled"])
        self.search_btn.state(["!disabled"])
        msg = f"Read {len(self.read_entries)} file{'s' if len(self.read_entries) != 1 else ''}, {rows:,} rows."
        if failed:
            msg += f" {failed} could not be read (see the Size / rows column)."
        self.status.set(msg)
        self._refresh_tree()

    def _build_header_options(self):
        hmap = {}
        for e in self.read_entries:
            for s in e.data:
                for h in s.headers:
                    k = h.lower()
                    hmap.setdefault(k, [h, set()])[1].add(e.iid)
        items = sorted(hmap.items(), key=lambda kv: (-len(kv[1][1]), kv[1][0].lower()))
        self.idcol_map = {}
        labels = []
        for k, (label, files) in items:
            disp = f"{label}  ({len(files)} file{'s' if len(files) != 1 else ''})"
            self.idcol_map[disp] = k
            labels.append(disp)
        self.idbox["values"] = labels
        prev_key = self.idcol_map.get(self.id_col.get())
        keys = [k for k, _ in items]
        guess = None
        if prev_key in keys:
            guess = prev_key
        else:
            guess = (next((k for k in keys if k == "id"), None)
                     or next((k for k in keys if re.search(r"(^|[^a-z])id($|[^a-z])", k)), None)
                     or next((k for k in keys if "id" in k), None)
                     or (keys[0] if keys else None))
        for disp, k in self.idcol_map.items():
            if k == guess:
                self.id_col.set(disp)
                break
        else:
            self.id_col.set("")

    # ---------------- search inputs ----------------
    def _on_mode(self):
        if self.mode.get() == "id":
            self.idrow.grid()
            self.terms_label.configure(text="IDs to find (one per line, or separated by commas)")
        else:
            self.idrow.grid_remove()
            self.terms_label.configure(text="Values to find (a row matches if any cell contains one)")

    def _choose_bulk(self):
        p = filedialog.askopenfilename(
            title="Choose a list file (first column is used)",
            filetypes=[("Text and Excel", "*.txt *.csv *.tsv *.xlsx *.xlsm *.xls"), ("All files", "*.*")])
        if p:
            self.bulk_path = p
            self._reload_bulk()

    def _reload_bulk(self):
        if not self.bulk_path:
            return
        try:
            self.bulk = read_first_column(self.bulk_path, self.bulk_header.get())
        except Exception as ex:
            messagebox.showerror("Folder lookup", f"Could not read the list file.\n\n{ex}")
            return
        self.bulk_text.set(f"{os.path.basename(self.bulk_path)}: {len(self.bulk):,} values")
        self.bulk_row.grid()
        self.msg.set("")

    def _clear_bulk(self):
        self.bulk, self.bulk_path = [], None
        self.bulk_row.grid_remove()

    # ---------------- searching ----------------
    def _search_clicked(self):
        self.msg.set("")
        terms = list(dict.fromkeys(parse_terms(self.terms.get("1.0", "end")) + self.bulk))
        if not terms:
            self.msg.set("Add at least one value, or upload a list.")
            return
        if not self.loaded:
            self.msg.set("Read the files first.")
            return
        mode = self.mode.get()
        id_key = self.idcol_map.get(self.id_col.get(), "")
        if mode == "id" and not id_key:
            self.msg.set("Choose an ID header.")
            return
        entries = list(self.read_entries)
        self.phase = "Searching"
        self.search_btn.state(["disabled"])
        self.progress.configure(maximum=max(1, len(entries)), value=0)
        self.run_bg(lambda prog: search(entries, mode, id_key, terms, prog), self._search_done)

    def _search_done(self, res):
        self.results = res
        self.search_btn.state(["!disabled"])
        groups = res["groups"]
        total = sum(len(g["rows"]) for g in groups)
        files = len({g["path"] for g in groups})
        found = res["term_count"] - len(res["not_found"])

        # table with the union of headers
        keys, labels = [], {}
        for g in groups:
            for h in g["headers"]:
                k = h.lower()
                if k not in labels:
                    labels[k] = h
                    keys.append(k)
        cols = [f"c{i}" for i in range(4 + len(keys))]
        self.rtree.delete(*self.rtree.get_children())
        self.rtree.configure(columns=cols)
        heads = ["File", "Sheet", "Row", "Matched"] + [labels[k] for k in keys]
        for c, h in zip(cols, heads):
            self.rtree.heading(c, text=h)
            self.rtree.column(c, width=150 if c != "c2" else 60, stretch=False)
        shown = 0
        for g in groups:
            idx = [next((i for i, h in enumerate(g["headers"]) if h.lower() == k), -1) for k in keys]
            for n, cells, term in g["rows"]:
                if shown >= MAX_SHOWN:
                    break
                vals = [g["path"], g["sheet"] or "", n, term] + [
                    cells[i] if 0 <= i < len(cells) else "" for i in idx]
                self.rtree.insert("", "end", values=vals)
                shown += 1

        s = (f"{total:,} row{'s' if total != 1 else ''} matched in {files} file{'s' if files != 1 else ''}. "
             f"{found:,} of {res['term_count']:,} value{'s' if res['term_count'] != 1 else ''} found.")
        if total > shown:
            s += f"  (Showing the first {shown:,}; exports include all rows.)"
        self.summary.set(s)
        self._set_text(self.nf_text, "\n".join(res["not_found"]) or "Every value was found.")
        self._set_text(self.skip_text, "\n".join(res["skipped"]) or "No files were skipped.")
        state = "!disabled" if total else "disabled"
        self.csv_btn.state([state])
        self.xlsx_btn.state([state])
        self.status.set(self.status.get().split("  ")[0])

    # ---------------- export ----------------
    def _export_csv(self):
        p = filedialog.asksaveasfilename(defaultextension=".csv", initialfile="lookup-results.csv",
                                         filetypes=[("CSV", "*.csv")])
        if not p:
            return
        try:
            with open(p, "w", newline="", encoding="utf-8-sig") as f:
                csv.writer(f).writerows(export_table(self.results["groups"]))
        except OSError as ex:
            messagebox.showerror("Folder lookup", f"Could not save the file.\n\n{ex}")
            return
        messagebox.showinfo("Folder lookup", f"Saved {p}")

    def _export_xlsx(self):
        if openpyxl is None:
            messagebox.showerror("Folder lookup", "Excel export needs:  pip install openpyxl")
            return
        p = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="lookup-results.xlsx",
                                         filetypes=[("Excel", "*.xlsx")])
        if not p:
            return
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Results"
            for row in export_table(self.results["groups"]):
                ws.append([ILLEGAL_XLSX.sub("", str(c)) for c in row])
            wb.save(p)
        except Exception as ex:
            messagebox.showerror("Folder lookup", f"Could not save the file.\n\n{ex}")
            return
        messagebox.showinfo("Folder lookup", f"Saved {p}")


def main():
    try:  # crisp text on high-DPI Windows screens
        import ctypes
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass
    App().mainloop()


if __name__ == "__main__":
    main()
